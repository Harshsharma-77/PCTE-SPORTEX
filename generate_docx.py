import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_synopsis_docx(output_path):
    doc = docx.Document()

    # Set Margins (Standard 1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styling Palettes
    COLOR_PRIMARY = RGBColor(15, 76, 129)     # Deep Royal Navy
    COLOR_SECONDARY = RGBColor(40, 116, 166)  # Blue Slate
    COLOR_TEXT = RGBColor(44, 62, 80)         # Charcoal Black

    # Configure Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # COVER / HEADER TITLE BLOCK
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("A MAJOR PROJECT SYNOPSIS ON\n")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    run_title = title_p.add_run("PCTE SPORTEX\nCOLLEGE SPORTS MANAGEMENT & REAL-TIME LIVE SCORING PLATFORM\n")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    inst_p = doc.add_paragraph()
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst_p.paragraph_format.space_after = Pt(20)
    r_inst = inst_p.add_run(
        "Submitted in partial fulfillment of the requirements for the degree of\n"
        "BACHELOR OF TECHNOLOGY\n"
        "IN\n"
        "COMPUTER SCIENCE & ENGINEERING\n\n"
        "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n"
        "PUNJAB COLLEGE OF TECHNICAL EDUCATION (PCTE), LUDHIANA\n"
        "ACADEMIC SESSION: 2025 – 2026\n"
    )
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph("―" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Helper function for Section Headings
    def add_section_heading(number_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r_num = h.add_run(f"{number_str} ")
        r_num.font.bold = True
        r_num.font.size = Pt(14)
        r_num.font.color.rgb = COLOR_PRIMARY
        r_txt = h.add_run(title_str.upper())
        r_txt.font.bold = True
        r_txt.font.size = Pt(14)
        r_txt.font.color.rgb = COLOR_PRIMARY
        return h

    # Helper function for Subheadings
    def add_subheading(sub_str):
        sh = doc.add_paragraph()
        sh.paragraph_format.space_before = Pt(10)
        sh.paragraph_format.space_after = Pt(4)
        sh.paragraph_format.keep_with_next = True
        r = sh.add_run(sub_str)
        r.font.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_SECONDARY
        return sh

    # -------------------------------------------------------------
    # 1. INTRODUCTION
    # -------------------------------------------------------------
    add_section_heading("1.", "Introduction & Background")
    p1 = doc.add_paragraph(
        "Collegiate athletic programs serve as a cornerstone of physical development, institutional pride, "
        "and active student life. However, modern collegiate sports administration across universities in India "
        "is predominantly managed via disjointed spreadsheets, physical paper scorecards, and physical notice board postings. "
        "This fragmented operational method causes fixture delays, scoring inconsistencies, and unverified participation certificates."
    )
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p2 = doc.add_paragraph(
        "PCTE SPORTEX is an enterprise-grade, cloud-native Software-as-a-Service (SaaS) application engineered to provide a "
        "unified digital sports ecosystem. Developed from the ground up using modern full-stack web technologies (TypeScript, "
        "React 18, Node.js Express, PostgreSQL, Prisma ORM, and Socket.IO), SPORTEX seamlessly coordinates university sports directors, "
        "departmental team coaches, certified referees, and student athletes within a single responsive portal."
    )
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # -------------------------------------------------------------
    # 2. PROBLEM STATEMENT
    # -------------------------------------------------------------
    add_section_heading("2.", "Problem Statement")
    doc.add_paragraph("The current manual and semi-digital methods of collegiate sports governance suffer from critical challenges:")
    
    problems = [
        ("Manual Scheduling Collisions: ", "Venues and grounds are vulnerable to overlapping schedules and double-booking due to lack of automated conflict detection."),
        ("Absence of Live Telemetry: ", "Students, faculty, and alumni unable to physically attend campus grounds have no access to real-time, zero-delay match score updates."),
        ("Delayed Tournament Standings: ", "Points tables, goal differences, and leaderboards require laborious manual tallying, often taking days after match completion."),
        ("Fraud-Prone Paper Certificates: ", "Traditional paper athletic certificates can easily be duplicated, altered, or forged without instant verification."),
        ("Equipment Inventory Loss: ", "Sports goods and athletic equipment checkout processes lack real-time stock and condition inspection tracking.")
    ]
    for bold_prefix, text in problems:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(bold_prefix)
        r_b.font.bold = True
        p.add_run(text)

    # -------------------------------------------------------------
    # 3. OBJECTIVES
    # -------------------------------------------------------------
    add_section_heading("3.", "Objectives of the Project")
    doc.add_paragraph("The primary objectives of PCTE SPORTEX are as follows:")
    
    objectives = [
        "Implement Granular Role-Based Access Control (RBAC) across 6 roles (Super Admin, Sports Admin, Coach, Referee, Student Athlete, and Viewer).",
        "Engineer a Real-Time Live Scoring Engine using Socket.IO WebSockets delivering game clock control, sport-specific events (goals, runs, points), and commentary with <50ms latency.",
        "Provide an Automated Tournament & Bracket Pairing Engine supporting Single Elimination Knockouts and Round-Robin leagues with automated post-match points cascade.",
        "Deploy a Cryptographic QR Certificate Engine with UUID verification tokens allowing instant public zero-authentication verification against the database.",
        "Build an Intelligent Ground Conflict Detection algorithm that prevents overlapping fixtures at the same sports arena.",
        "Incorporate an Equipment Inventory Desk with condition auditing (New, Good, Fair, Damaged) and loan ledger.",
        "Deliver Global Omnisearch (Ctrl+K) for instant query resolution across athletes, teams, sports, fixtures, and announcements."
    ]
    for obj in objectives:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(obj)

    # -------------------------------------------------------------
    # 4. SYSTEM ARCHITECTURE & METHODOLOGY
    # -------------------------------------------------------------
    add_section_heading("4.", "System Architecture & Methodology")
    p_arch = doc.add_paragraph(
        "PCTE SPORTEX follows a modern Three-Tier Client-Server Architecture, dividing concerns into presentation, "
        "business logic, and persistence layers for maximum reliability, speed, and horizontal scalability."
    )
    p_arch.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    arch_layers = [
        ("Presentation Layer (Frontend): ", "Built using React 18, TypeScript, Tailwind CSS, shadcn/ui components, and Framer Motion. Bundled and deployed on Vercel."),
        ("Application & Streaming Layer (Backend): ", "Engineered with Node.js, Express 4, TypeScript, Zod schema validation, Helmet security, and Socket.IO. Deployed on Render."),
        ("Persistence Layer (Database): ", "PostgreSQL 16+ database managed via Prisma ORM with 24 normalized relational entities and connection pooling. Hosted on Neon Serverless Cloud.")
    ]
    for layer, desc in arch_layers:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(layer)
        r_b.font.bold = True
        p.add_run(desc)

    # -------------------------------------------------------------
    # 5. FUNCTIONAL MODULES TABLE
    # -------------------------------------------------------------
    add_section_heading("5.", "Functional Modules & Capabilities")
    
    table_modules = doc.add_table(rows=1, cols=3)
    table_modules.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_modules.autofit = False

    hdr_cells = table_modules.rows[0].cells
    hdr_cells[0].text = "Module Name"
    hdr_cells[1].text = "Primary User"
    hdr_cells[2].text = "Key Functionalities & Features"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shading = parse_xml(r'<w:shd {} w:fill="0F4C81"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    modules_data = [
        ("Authentication & RBAC", "All Users", "Dual-token JWT, bcrypt password hashing, 6-role permission matrix, session protection."),
        ("Sports & Franchise Directory", "Admin, Coach", "Configure sports disciplines, departmental teams, squads, jersey numbers, and captains."),
        ("Tournaments & Scheduling", "Sports Admin", "Knockout/Round-robin brackets, automated fixture generation, venue conflict detector."),
        ("Live Match Scoring Console", "Match Referee", "WebSocket real-time scoreboard, timer kickoff/pause, goals, cards, and text commentary."),
        ("Standings & Leaderboards", "Public / All", "Automated points table recalculation (GD, Points, Form Guide), top scorers, and stats."),
        ("Verifiable QR Certificates", "Student / Admin", "Institutional certificate issuing engine, cryptographic UUID token, public QR verification."),
        ("Equipment Inventory Desk", "Admin, Coach", "Item tracking, condition grades (New/Good/Damaged), loan checkout, fine calculations."),
        ("Grievance & Audit Trail", "Super Admin", "Student dispute resolution desk and immutable audit ledger of all system transactions.")
    ]

    for mod, user, feat in modules_data:
        row_cells = table_modules.add_row().cells
        row_cells[0].text = mod
        row_cells[1].text = user
        row_cells[2].text = feat
        for cell in row_cells:
            cell.paragraphs[0].runs[0].font.size = Pt(10)

    # Set column widths
    for row in table_modules.rows:
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(1.2)
        row.cells[2].width = Inches(3.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 6. HARDWARE & SOFTWARE SPECIFICATIONS
    # -------------------------------------------------------------
    add_section_heading("6.", "System Specifications")
    add_subheading("Software Requirements")

    table_sw = doc.add_table(rows=1, cols=2)
    table_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    sw_hdr = table_sw.rows[0].cells
    sw_hdr[0].text = "Component"
    sw_hdr[1].text = "Technology / Tool Utilized"
    for cell in sw_hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shading = parse_xml(r'<w:shd {} w:fill="2874A6"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    sw_specs = [
        ("Frontend Technology", "React 18, TypeScript, Tailwind CSS, Vite"),
        ("Backend Framework", "Node.js (v18+), Express 4, TypeScript"),
        ("Database Engine", "PostgreSQL 16+ (Serverless Neon Cloud)"),
        ("Object-Relational Mapping", "Prisma ORM (v5.22.0)"),
        ("Real-Time Communications", "Socket.IO (v4.8.1) WebSocket Protocol"),
        ("Security Standards", "JWT (HS256), BcryptJS, Helmet, Zod Validation"),
        ("Deployment Cloud Hosting", "Frontend on Vercel, Backend API on Render")
    ]
    for comp, tech in sw_specs:
        row = table_sw.add_row().cells
        row[0].text = comp
        row[1].text = tech
        row[0].paragraphs[0].runs[0].font.size = Pt(10)
        row[1].paragraphs[0].runs[0].font.size = Pt(10)
    for row in table_sw.rows:
        row.cells[0].width = Inches(2.5)
        row.cells[1].width = Inches(4.0)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_subheading("Hardware Requirements")
    hw_list = [
        "Processor: Intel Core i3 / AMD Ryzen 3 or higher.",
        "Memory (RAM): Minimum 4 GB (8 GB recommended).",
        "Storage: 500 MB minimum available hard disk space.",
        "Network: Broadband / Wi-Fi internet connectivity for cloud database and WebSocket streaming."
    ]
    for hw in hw_list:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(hw)

    # -------------------------------------------------------------
    # 7. LIVE DEPLOYMENT & VERIFICATION
    # -------------------------------------------------------------
    add_section_heading("7.", "Live Deployment & Verification Details")
    p_dep = doc.add_paragraph(
        "PCTE SPORTEX has been fully built, verified, and successfully deployed to live cloud infrastructure:"
    )
    
    dep_items = [
        ("Live Web Application: ", "https://pcte-sportex.vercel.app"),
        ("Production API Gateway: ", "https://pcte-sportex-api.onrender.com"),
        ("Cloud Database Instance: ", "Serverless PostgreSQL on Neon.tech (US-East Region)"),
        ("Source Code Repository: ", "https://github.com/Harshsharma-77/PCTE-SPORTEX")
    ]
    for title, link in dep_items:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(title)
        r_b.font.bold = True
        r_l = p.add_run(link)
        r_l.font.color.rgb = COLOR_SECONDARY
        r_l.font.underline = True

    # -------------------------------------------------------------
    # 8. EXPECTED OUTCOMES & CONCLUSION
    # -------------------------------------------------------------
    add_section_heading("8.", "Expected Outcomes & Conclusion")
    p_conc = doc.add_paragraph(
        "PCTE SPORTEX transforms university sports administration from an error-prone, paper-heavy routine into an "
        "integrated, real-time digital experience. By unifying tournament brackets, real-time WebSocket scoring, ground conflict "
        "prevention, and cryptographic certificate verification, SPORTEX sets a benchmark for collegiate athletic software systems. "
        "The project stands 100% complete, fully verified, and ready for institutional deployment across collegiate championships."
    )
    p_conc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # -------------------------------------------------------------
    # SIGNATURE BLOCK
    # -------------------------------------------------------------
    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2 = sig_table.rows[0].cells
    c1.width = Inches(3.25)
    c2.width = Inches(3.25)

    p_sig1 = c1.paragraphs[0]
    p_sig1.add_run("Submitted By:\n\n\n_______________________\nHarsh Sharma\nUniversity Roll No: 2329263\nB.Tech CSE, 6th Semester").font.size = Pt(11)

    p_sig2 = c2.paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.add_run("Approved By (Project Guide):\n\n\n_______________________\nDepartment of CSE\nPCTE, Ludhiana").font.size = Pt(11)

    # Save Document
    doc.save(output_path)
    print(f"Successfully generated Word Document at: {output_path}")

if __name__ == '__main__':
    create_synopsis_docx("c:\\Users\\Admin\\Desktop\\AQUA\\pcte-sportex\\PCTE_SPORTEX_Project_Synopsis.docx")
